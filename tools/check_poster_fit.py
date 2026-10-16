"""Estimate whether any poster text box overflows its geometry.

    python tools/check_poster_fit.py

PowerPoint is not available here, so this approximates the wrapped line count
from the box width, the run font size and an average glyph advance, and
compares it with the number of lines the box can hold. It is a guard against
silently shipping a poster with clipped text, not a substitute for opening the
file. Anything reported TIGHT or OVERFLOW should be re-checked by eye in
PowerPoint before exporting the PDF.
"""
from pathlib import Path
import math
import sys

from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
POSTER = ROOT / "poster" / "final-poster.pptx"
MM = 914400 / 25.4

# average glyph advance and line pitch as a fraction of the point size
GLYPH, PITCH, INSET_W, INSET_H = 0.50, 1.22, 6.0, 3.0


def main():
    prs = Presentation(str(POSTER))
    slide = prs.slides[0]
    w_mm, h_mm = prs.slide_width / MM, prs.slide_height / MM
    print(f"slide {w_mm:.0f} x {h_mm:.0f} mm "
          f"({'A0 landscape OK' if abs(w_mm-1189)<1 and abs(h_mm-841)<1 else 'NOT A0'})")
    print(f"\n{'shape':26s}{'W':>5s}{'H':>5s}{'pt':>6s}{'fill':>7s}  status")

    worst = []
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        txt = sh.text_frame.text
        if len(txt.strip()) < 40:
            continue
        sizes = [r.font.size.pt for p in sh.text_frame.paragraphs
                 for r in p.runs if r.font.size]
        pt = sizes[0] if sizes else 12.0
        bw, bh = sh.width / MM, sh.height / MM
        cols = max(1.0, (bw - INSET_W) / (GLYPH * pt * 25.4 / 72))
        rows = max(1.0, (bh - INSET_H) / (PITCH * pt * 25.4 / 72))
        lines = sum(max(1, math.ceil(len(par) / cols)) for par in txt.split("\n"))
        fill = lines / rows
        status = "OK" if fill <= 1.0 else ("TIGHT" if fill < 1.15 else "OVERFLOW")
        if status != "OK":
            worst.append((sh.name, fill, status))
        print(f"{sh.name:26s}{bw:5.0f}{bh:5.0f}{pt:6.1f}{fill:7.2f}  {status}")

    # shapes must stay on the canvas
    off = [sh.name for sh in slide.shapes
           if sh.left is not None and sh.width is not None
           and (sh.left / MM < -1 or (sh.left + sh.width) / MM > w_mm + 1
                or sh.top / MM < -1 or (sh.top + sh.height) / MM > h_mm + 1)]
    print("\noff-canvas shapes:", ", ".join(off) if off else "none")
    if worst:
        print("\nNEEDS ATTENTION:")
        for n, f, st in sorted(worst, key=lambda x: -x[1]):
            print(f"  {n:26s} fill {f:.2f}  {st}")
        return 1
    print("\nAll text boxes fit at the estimated metrics.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
